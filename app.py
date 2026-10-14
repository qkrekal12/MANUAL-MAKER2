from __future__ import annotations

import base64
import hashlib
import html
import io
import os
import re
from pathlib import Path

import streamlit as st
from openai import OpenAI
from PIL import Image, ImageDraw, ImageFont
from streamlit_image_coordinates import streamlit_image_coordinates


st.set_page_config(page_title="pagewise | 매뉴얼 스튜디오", page_icon="📘", layout="wide")
st.markdown("""
<style>
:root {--pw-purple:#6654e8;--pw-ink:#202b43;--pw-muted:#8290ad;--pw-line:#e5e9f1;--pw-bg:#f5f6fa;}
.stApp {background:var(--pw-bg); color:var(--pw-ink);}
[data-testid="stHeader"] {background:transparent;}
[data-testid="stSidebar"] {background:#fff;border-right:1px solid var(--pw-line);min-width:265px!important;}
[data-testid="stSidebar"] > div:first-child {padding:18px 16px 28px;}
.block-container {max-width:1120px;padding:25px 28px 4rem;}
[data-testid="stTextInput"] input,[data-testid="stTextArea"] textarea,[data-testid="stSelectbox"] [role="combobox"] {border-radius:8px;border-color:#dfe4ed;background:#fff;}
[data-testid="stFileUploader"] section {border:1px dashed #cfd5e3;border-radius:12px;background:#fafbfe;}
[data-testid="stVerticalBlockBorderWrapper"] > div {border:1px solid var(--pw-line);border-radius:14px;background:#fff;}
.stButton button,.stDownloadButton button {border-radius:9px;border-color:#e1e5ee;color:#4e5a73;font-weight:600;}
.stButton button[kind="primary"] {background:var(--pw-purple);border-color:var(--pw-purple);color:#fff;}
.pw-brand {font-size:23px;font-weight:800;letter-spacing:-1px;color:#202b43;margin:8px 4px 38px;}
.pw-brand b {display:inline-grid;place-items:center;width:34px;height:34px;border-radius:11px;background:#6654e8;color:white;margin-right:8px;}
.pw-side-label {color:#9aa3b5;font-size:11px;font-weight:800;letter-spacing:1.3px;margin:22px 8px 9px;}
.pw-workspace {border:1px solid #e7eaf1;border-radius:13px;padding:13px 10px;background:#fff;font-size:13px;}
.pw-workspace small,.pw-muted {color:var(--pw-muted);font-size:12px;}
.pw-side-help {background:#f4f3ff;border:1px solid #eeebff;border-radius:14px;padding:16px;margin-top:100px;font-size:13px;line-height:1.65;}
.pw-topline {display:flex;align-items:center;justify-content:space-between;border-bottom:1px solid var(--pw-line);padding:0 0 16px;margin-bottom:34px;color:#8490a9;font-size:13px;}
.pw-eyebrow {font-size:11px;font-weight:800;letter-spacing:2px;color:#7868ed;text-transform:uppercase;margin-bottom:10px;}
.pw-hero {display:flex;align-items:center;justify-content:space-between;gap:22px;margin-bottom:23px;}
.pw-hero h1 {font-size:34px;line-height:1.18;letter-spacing:-1.3px;margin:0;color:#26334f;}
.pw-hero p {color:#8792aa;font-size:14px;margin:9px 0 0;}
.pw-pill {white-space:nowrap;background:#fff;border:1px solid var(--pw-line);border-radius:24px;padding:12px 17px;color:#77839e;font-size:12px;}
.pw-stepper {display:flex;align-items:center;gap:14px;background:#fff;border:1px solid var(--pw-line);border-radius:14px;padding:18px 22px;margin:0 0 18px;}
.pw-step {display:flex;align-items:center;gap:10px;color:#8a96ad;white-space:nowrap;font-size:12px;}
.pw-step i {display:grid;place-items:center;width:30px;height:30px;border-radius:50%;background:#f0efff;color:#6654e8;font-style:normal;font-size:10px;font-weight:700;}
.pw-step.active {color:#6654e8;font-weight:700;}.pw-step-line {height:1px;background:#e4e7ef;flex:1;}
.pw-doc {display:flex;align-items:center;justify-content:space-between;background:#fff;border:1px solid var(--pw-line);border-radius:14px;padding:16px 20px;margin-bottom:17px;}
.pw-doc-left {display:flex;align-items:center;gap:13px;}.pw-doc-icon {display:grid;place-items:center;width:44px;height:44px;border-radius:13px;background:#f0efff;color:#6654e8;font-size:20px;}
.pw-doc small {display:block;color:#8792aa;font-size:11px;margin-bottom:2px;}.pw-doc b {font-size:15px;color:#27354e;}
.pw-dot {display:inline-block;width:7px;height:7px;border-radius:50%;background:#55c69b;margin-right:7px;}
.pw-section {display:flex;align-items:center;gap:12px;margin:4px 0 17px;}.pw-section-no {display:grid;place-items:center;width:32px;height:32px;border-radius:50%;background:#f0efff;color:#6654e8;font-size:11px;font-weight:700;}
.pw-section h3 {font-size:17px;margin:0;color:#283650;}.pw-section p {font-size:12px;color:#8a96ad;margin:3px 0 0;}
.pw-upload-heading {margin-top:24px!important;}
.pw-privacy {display:flex;gap:10px;padding:14px 16px;background:#f5f7f4;border-radius:12px;color:#607087;font-size:12px;line-height:1.6;margin:13px 0;}
.pw-screen-title {font-size:15px;font-weight:700;color:#34415c;margin:16px 0 8px;}
.pw-result-title {display:flex;align-items:center;justify-content:space-between;margin:34px 0 16px;}
.pw-result-title h2 {font-size:21px;margin:0;color:#26334f;}
@media(max-width:850px){.block-container{padding:18px 16px 3rem}.pw-hero{align-items:flex-start;flex-direction:column}.pw-pill{white-space:normal}.pw-stepper{overflow:auto}.pw-step-line{min-width:24px}}
</style>
""", unsafe_allow_html=True)


def api_key() -> str:
    key = os.getenv("OPENAI_API_KEY", "").strip()
    if key:
        return key
    try:
        key = str(st.secrets.get("OPENAI_API_KEY", "")).strip()
        if key:
            return key
    except Exception:
        # Streamlit Cloud secrets may not be configured yet; local .env is also supported.
        pass
    for parent in (Path(__file__).resolve().parent, *Path(__file__).resolve().parents):
        env_file = parent / ".env"
        if env_file.is_file():
            try:
                for line in env_file.read_text(encoding="utf-8").splitlines():
                    if line.strip().startswith("OPENAI_API_KEY="):
                        return line.split("=", 1)[1].strip().strip('"').strip("'")
            except OSError:
                pass
    return ""


def natural_key(name: str):
    return [(0, int(part)) if part.isdigit() else (1, part.casefold()) for part in re.split(r"(\d+)", name)]


def marker_image(image: Image.Image, markers: list[dict]) -> bytes:
    result = image.convert("RGB").copy()
    draw = ImageDraw.Draw(result)
    radius = max(15, min(result.width, result.height) // 32)
    try:
        font = ImageFont.truetype("malgun.ttf", radius)
    except OSError:
        font = ImageFont.load_default()
    for idx, marker in enumerate(markers, 1):
        x = int(marker["x"] * result.width)
        y = int(marker["y"] * result.height)
        draw.ellipse((x-radius, y-radius, x+radius, y+radius), fill="#6656e8", outline="white", width=3)
        label = str(idx)
        box = draw.textbbox((0, 0), label, font=font)
        draw.text((x-(box[2]-box[0])/2, y-(box[3]-box[1])/2-box[1]), label, fill="white", font=font)
    output = io.BytesIO()
    result.save(output, format="PNG")
    return output.getvalue()


def image_as_data(data: bytes) -> str:
    return "data:image/png;base64," + base64.b64encode(data).decode("ascii")


def make_html(markdown: str, screens: list[dict], title: str) -> bytes:
    import markdown as md
    body = "".join(f'<section class="screen"><img src="{image_as_data(s["png"])}"><div><h3>STEP {n}</h3><p>{html.escape(s["caption"])}</p>' + "".join(f'<p><b>{i}. {html.escape(m["label"] or "버튼 설명")}</b><br>{html.escape(m["description"] or "설명을 입력하세요.")}</p>' for i, m in enumerate(s["markers"], 1)) + "</div></section>" for n, s in enumerate(screens, 1))
    safe_title = html.escape(title)
    formatted_manual = md.markdown(markdown, extensions=["tables", "fenced_code"])
    text = f"""<!doctype html><html lang="ko"><meta charset="utf-8"><meta name="viewport" content="width=device-width"><title>{safe_title}</title><style>
    body{{font-family:Arial,'Malgun Gothic',sans-serif;color:#263248;max-width:1080px;margin:40px auto;padding:0 28px;line-height:1.75}}h1,h2,h3{{color:#263b68}}.manual{{white-space:normal}}.screen{{display:grid;grid-template-columns:minmax(0,1.25fr) minmax(250px,.75fr);gap:24px;border:1px solid #dce2ed;border-radius:14px;padding:18px;margin:22px 0;align-items:start}}.screen img{{width:100%;height:auto;border-radius:8px}}.screen p{{margin:.5em 0}}@media(max-width:700px){{.screen{{grid-template-columns:1fr}}}}@media print{{body{{margin:15mm auto}}.screen{{break-inside:avoid}}}}
    </style><h1>{safe_title}</h1><article class="manual">{formatted_manual}</article><h2>화면별 안내</h2>{body}</html>"""
    return text.encode("utf-8")


def make_docx(markdown: str, screens: list[dict], title: str) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
        section.left_margin = Inches(.7); section.right_margin = Inches(.7)
    doc.add_heading(title, 0)
    lines = markdown.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not cells or all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    index += 1
                    continue
                rows.append(cells); index += 1
            if rows:
                width = max(map(len, rows))
                table = doc.add_table(rows=0, cols=width); table.style = "Table Grid"
                for row_no, values in enumerate(rows):
                    cells = table.add_row().cells
                    for col_no in range(width):
                        cell_text = values[col_no] if col_no < len(values) else ""
                        paragraph = cells[col_no].paragraphs[0]
                        for part_no, part in enumerate(cell_text.split("**")):
                            if part:
                                run = paragraph.add_run(part); run.bold = row_no == 0 or part_no % 2 == 1
                doc.add_paragraph()
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith(("- ", "* ")):
            doc.add_paragraph(line[2:].replace("**", ""), style="List Bullet")
        elif re.match(r"^\d+\.\s", line):
            doc.add_paragraph(re.sub(r"^\d+\.\s", "", line).replace("**", ""), style="List Number")
        else:
            p = doc.add_paragraph()
            for part_no, part in enumerate(line.split("**")):
                if part:
                    run = p.add_run(part); run.bold = part_no % 2 == 1
        index += 1
    if screens:
        doc.add_heading("화면별 안내", 1)
    for n, screen in enumerate(screens, 1):
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"; table.alignment = WD_TABLE_ALIGNMENT.CENTER
        left, right = table.rows[0].cells
        pic = left.paragraphs[0]; pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.add_run().add_picture(io.BytesIO(screen["png"]), width=Inches(3.7))
        right.paragraphs[0].add_run(f"STEP {n} · {screen['caption']}").bold = True
        for i, marker in enumerate(screen["markers"], 1):
            p = right.add_paragraph()
            p.add_run(f"{i}. {marker['label'] or '버튼 설명'}: ").bold = True
            p.add_run(marker["description"] or "설명을 입력하세요.")
        doc.add_paragraph()
    out = io.BytesIO(); doc.save(out); return out.getvalue()


def make_pdf(markdown: str, screens: list[dict], title: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as PDFImage, Table, TableStyle, PageBreak

    font_path = next((p for p in (Path(r"C:\Windows\Fonts\malgun.ttf"), Path(r"C:\Windows\Fonts\gulim.ttc")) if p.exists()), None)
    if font_path:
        pdfmetrics.registerFont(TTFont("Korean", str(font_path), subfontIndex=0))
        font_name = "Korean"
    else:
        font_name = "Helvetica"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="KBody", parent=styles["BodyText"], fontName=font_name, fontSize=9.5, leading=15, spaceAfter=6))
    styles.add(ParagraphStyle(name="KTitle", parent=styles["Title"], fontName=font_name, fontSize=20, leading=26, alignment=TA_LEFT, textColor=colors.HexColor("#263b68")))
    styles.add(ParagraphStyle(name="KHead", parent=styles["Heading2"], fontName=font_name, fontSize=13, leading=18, textColor=colors.HexColor("#263b68")))
    story = [Paragraph(html.escape(title), styles["KTitle"]), Spacer(1, 8)]
    for raw in markdown.splitlines():
        line = raw.strip()
        if not line:
            continue
        if line.startswith("## "):
            story.append(Paragraph(html.escape(line[3:]), styles["KHead"]))
        elif line.startswith("### "):
            story.append(Paragraph(html.escape(line[4:]), styles["KHead"]))
        else:
            story.append(Paragraph(html.escape(line.lstrip("# ")).replace("**", ""), styles["KBody"]))
    for n, screen in enumerate(screens, 1):
        story.append(Paragraph(f"STEP {n} · {html.escape(screen['caption'])}", styles["KHead"]))
        picture = PDFImage(io.BytesIO(screen["png"]), width=105*mm, height=65*mm, kind="proportional")
        details = [Paragraph(html.escape(screen["caption"]), styles["KBody"])]
        details.extend(Paragraph(f"{i}. {html.escape(m['label'] or '버튼 설명')}<br/>{html.escape(m['description'] or '설명을 입력하세요.')}", styles["KBody"]) for i, m in enumerate(screen["markers"], 1))
        table = Table([[picture, details]], colWidths=[110*mm, 75*mm])
        table.setStyle(TableStyle([("VALIGN", (0,0), (-1,-1), "TOP"), ("BOX", (0,0), (-1,-1), .5, colors.HexColor("#dce2ed")), ("LEFTPADDING", (0,0), (-1,-1), 6), ("RIGHTPADDING", (0,0), (-1,-1), 6)]))
        story.extend([table, Spacer(1, 10)])
    out = io.BytesIO(); SimpleDocTemplate(out, pagesize=A4, rightMargin=15*mm, leftMargin=15*mm, topMargin=15*mm, bottomMargin=15*mm).build(story)
    return out.getvalue()


with st.sidebar:
    st.markdown('<div class="pw-brand"><b>▤</b>pagewise.</div>', unsafe_allow_html=True)
    st.markdown('<div class="pw-side-label">WORKSPACE</div><div class="pw-workspace"><b>◉　로컬 워크스페이스</b><br><small>　　이 PC에서 작업 중</small></div><div class="pw-side-label">작업</div>', unsafe_allow_html=True)
    if st.button("✳　 새 매뉴얼 만들기", use_container_width=True, key="side_new"):
        st.session_state.current_page = "studio"
        st.rerun()
    if st.button("◷　 최근 작업", use_container_width=True, key="side_recent"):
        st.session_state.focus_result = True
    st.markdown('<div class="pw-side-help"><b>✦　첫 매뉴얼을 만들고 있나요?</b><br><span class="pw-muted">화면 이미지를 올리면 사용 설명 초안을 만들어요.</span><br><span style="color:#6654e8;font-weight:700">사용 방법 보기 ↗</span></div>', unsafe_allow_html=True)
    st.markdown('<hr><div class="pw-muted">●　매뉴얼 스튜디오<br>　　로컬 작업공간</div>', unsafe_allow_html=True)

if st.session_state.get("current_page") == "proposal":
    st.markdown('<div class="pw-topline"><span>워크스페이스　/　기획서</span><span>pagewise.</span></div>', unsafe_allow_html=True)
    if st.button("← 매뉴얼 만들기로 돌아가기", key="proposal_back"):
        st.session_state.current_page = "studio"
        st.rerun()
    proposal = Path(__file__).resolve().parent / "AX_PoC_기획서 (1).html"
    if proposal.is_file():
        st.components.v1.html(proposal.read_text(encoding="utf-8"), height=950, scrolling=True)
    else:
        st.info("기획서 파일이 배포본에 포함되지 않았습니다.")
    st.stop()

top_left, top_right1, top_right2, top_right3 = st.columns([5.4, 1.25, 1.1, .75])
with top_left:
    st.markdown('<div class="pw-topline" style="border:0;margin:0;padding:8px 0 0">워크스페이스　/　새 매뉴얼 만들기</div>', unsafe_allow_html=True)
with top_right1:
    if st.button("기획서 첫 페이지", use_container_width=True, key="open_proposal"):
        st.session_state.current_page = "proposal"
        st.rerun()
with top_right2:
    demo_clicked = st.button("데모 입력", use_container_width=True, key="fill_demo")
with top_right3:
    if st.button("새 작업", use_container_width=True, key="new_task"):
        for key in ["function_name", "service_name", "notes", "manual", "manual_screens", "manual_title", "screens"]:
            st.session_state.pop(key, None)
        st.rerun()

if demo_clicked:
    st.session_state.function_name = "신규입회"
    st.session_state.service_name = "사내 업무 포털"
    st.session_state.notes = "[데모 예시] 사용자가 입회 정보를 입력하고 필요한 자료를 첨부해 신청하는 기능을 설명해 주세요. 실제 화면과 완료 정책을 확인해 매뉴얼을 작성하세요."

st.markdown('<div class="pw-hero"><div><div class="pw-eyebrow">MANUAL STUDIO · AI ASSISTED</div><h1>화면에서 시작하는<br>사용자 매뉴얼</h1><p>화면 이미지와 개발 메모를 넣으면 사용 순서에 맞춰 매뉴얼을 작성합니다.</p></div><div class="pw-pill">분량과 구성은 기능에 맞게</div></div>', unsafe_allow_html=True)
st.markdown('<div class="pw-stepper"><div class="pw-step active"><i>01</i>화면 업로드</div><div class="pw-step-line"></div><div class="pw-step active"><i>02</i>화면 분석</div><div class="pw-step-line"></div><div class="pw-step"><i>03</i>내용 확인</div><div class="pw-step-line"></div><div class="pw-step"><i>04</i>내보내기</div></div>', unsafe_allow_html=True)
name_preview = html.escape(st.session_state.get("function_name", "").strip() or "새 사용자 매뉴얼")
count_preview = len(st.session_state.get("uploads", []))
st.markdown(f'<div class="pw-doc"><div class="pw-doc-left"><div class="pw-doc-icon">▤</div><div><small>매뉴얼 제목</small><b>{name_preview}</b></div></div><span class="pw-muted"><i class="pw-dot"></i>화면 {count_preview}개　│　자동 저장</span></div>', unsafe_allow_html=True)

st.markdown('<div class="pw-section"><div class="pw-section-no">01</div><div><h3>기능 정보</h3><p>화면을 정확히 이해하는 데 필요한 내용을 입력하세요.</p></div></div>', unsafe_allow_html=True)
with st.container(border=True):
    left_field, right_field = st.columns([1.35, 1])
    with left_field:
        name = st.text_input("기능 또는 서비스 이름", placeholder="예: 법인카드 신규 신청", key="function_name")
    with right_field:
        service = st.text_input("서비스명", placeholder="예: 사내 업무 포털", key="service_name")
    notes = st.text_area("개발 완료 내용 또는 추가 메모　선택", height=90, placeholder="이미지에 나오지 않는 처리 조건, 사용자 권한, 완료 결과, 알려진 오류 등을 입력하면 매뉴얼에 반영합니다.", key="notes")
    model = st.selectbox("분석 모델", ["gpt-6-luna"], format_func=lambda _: "GPT-6 Luna")

st.markdown('<div class="pw-section pw-upload-heading"><div class="pw-section-no">02</div><div><h3>화면 이미지 업로드　<span class="pw-muted">필수</span></h3><p>설명할 화면을 추가하세요. 여러 화면을 한 번에 올릴 수 있습니다.</p></div></div>', unsafe_allow_html=True)
uploads = st.file_uploader("화면 이미지 첨부 (PNG, JPG, JPEG, WEBP)", type=["png", "jpg", "jpeg", "webp"], accept_multiple_files=True, key="uploads", help="최대 10장, 이미지당 20MB 이하 권장")

if uploads:
    ordered = sorted(uploads, key=lambda f: natural_key(f.name))
    if len(ordered) > 10:
        st.warning("화면 이미지는 최대 10장까지 사용할 수 있습니다. 앞에서 10장만 반영합니다.")
        ordered = ordered[:10]
    too_large = [file.name for file in ordered if file.size > 20 * 1024 * 1024]
    if too_large:
        st.error("이미지당 20MB 이하로 줄인 뒤 다시 첨부해 주세요.")
        st.stop()
    if ordered != uploads:
        st.caption("파일명의 숫자를 기준으로 화면 순서를 정렬했습니다.")
    if "screens" not in st.session_state:
        st.session_state.screens = {}
    screens: list[dict] = []
    for step, uploaded in enumerate(ordered, 1):
        raw = uploaded.getvalue()
        fingerprint = hashlib.sha256(raw).hexdigest()[:16]
        state_key = f"screen_{fingerprint}"
        if state_key not in st.session_state.screens:
            st.session_state.screens[state_key] = {"markers": [], "last_event": None}
        state = st.session_state.screens[state_key]
        image = Image.open(io.BytesIO(raw)).convert("RGB")
        st.markdown(f'<div class="pw-screen-title">STEP {step}　화면 이미지 {step}</div>', unsafe_allow_html=True)
        col_image, col_info = st.columns([1.35, 1])
        with col_image:
            coords = streamlit_image_coordinates(image, width=760, key=f"click_{fingerprint}", cursor="crosshair")
            if coords and "x" in coords and "y" in coords:
                event = coords.get("time", (coords["x"], coords["y"]))
                if event != state["last_event"]:
                    state["markers"].append({"x": max(0, min(1, coords["x"] / image.width)), "y": max(0, min(1, coords["y"] / image.height)), "label": "", "description": ""})
                    state["last_event"] = event
                    st.rerun()
            if state["markers"]:
                st.image(marker_image(image, state["markers"]), caption="번호를 표시한 화면", use_container_width=True)
            st.caption("이미지를 클릭하면 해당 위치에 설명 번호가 추가됩니다.")
        with col_info:
            caption = st.text_input("화면 설명", key=f"caption_{fingerprint}", placeholder="이 화면에서 할 수 있는 일")
            if not state["markers"]:
                st.info("이미지에서 버튼이나 입력 위치를 클릭해 번호를 추가하세요. 선택 사항입니다.")
            for idx, marker in enumerate(state["markers"], 1):
                st.markdown(f"**{idx}번 설명**")
                marker["label"] = st.text_input("버튼/항목 이름", key=f"label_{fingerprint}_{idx}", value=marker["label"], placeholder="예: 조회")
                marker["description"] = st.text_area("설명", key=f"desc_{fingerprint}_{idx}", value=marker["description"], height=70, placeholder="이 항목을 선택하면 수행되는 동작")
                if st.button(f"{idx}번 삭제", key=f"del_{fingerprint}_{idx}"):
                    state["markers"].pop(idx - 1); st.rerun()
        screens.append({"image": image, "png": marker_image(image, state["markers"]), "caption": caption or f"화면 {step}의 주요 기능을 확인합니다.", "markers": state["markers"]})

    st.markdown('<div class="pw-privacy"><span>◈</span><span><b>개인정보를 확인하세요.</b> 화면 이미지와 메모는 매뉴얼 분석을 위해 OpenAI API로 전송됩니다. 민감 정보는 가린 뒤 첨부하세요.</span></div>', unsafe_allow_html=True)
    key = api_key()
    if not key:
        key = st.text_input("OpenAI API 키", type="password", help="입력한 키는 저장하지 않습니다.")
    st.caption("첨부 이미지는 분석 요청 처리 후 앱에서 파일로 보관하지 않습니다.")
    if st.button("화면 분석하고 매뉴얼 작성하기", type="primary", use_container_width=True, disabled=not (name.strip() and key.strip())):
        descriptions = []
        for i, screen in enumerate(screens, 1):
            desc = "; ".join(f"{j}. {m['label']}: {m['description']}" for j, m in enumerate(screen["markers"], 1) if m["label"] or m["description"])
            descriptions.append(f"STEP {i} 화면 설명: {screen['caption']}\n버튼/위치 설명: {desc or '사용자 제공 설명 없음'}")
        prompt = f"""IT 서비스 사용자를 위한 한국어 매뉴얼을 작성하세요. Markdown은 내부 구성용으로만 쓰며 사람에게 읽기 쉬운 문장과 표로 작성합니다.\n기능명: {name}\n서비스명: {service or '확인 필요'}\n참고 메모: {notes or '없음'}\n화면별 사용자 설명:\n{chr(10).join(descriptions)}\n\n기능에 맞는 목차와 장 수를 자율적으로 정하고, 실제 화면에 없는 정책·권한·오류·처리 결과를 추측하지 마세요. 불명확한 내용은 추가 확인 사항으로 따로 표시하세요. 단계별 절차를 사용자가 실제 화면에서 수행하는 순서로 쓰세요. 파일명을 노출하지 말고, 화면 이름은 STEP 번호로 표기하세요."""
        content = [{"type": "input_text", "text": prompt}]
        for screen in screens:
            content.append({"type": "input_image", "image_url": image_as_data(screen["png"]), "detail": "high"})
        try:
            with st.spinner("화면을 분석해 매뉴얼을 작성하고 있습니다..."):
                result = OpenAI(api_key=key.strip()).responses.create(model=model.strip(), input=[{"role": "user", "content": content}])
            st.session_state.manual = result.output_text
            st.session_state.manual_screens = screens
            st.session_state.manual_title = name.strip() + " 사용자 매뉴얼"
        except Exception as exc:
            st.error(f"매뉴얼을 작성하지 못했습니다: {exc}")

if st.session_state.get("manual"):
    markdown = st.session_state.manual
    saved_screens = st.session_state.manual_screens
    title = st.session_state.manual_title
    st.markdown('<div class="pw-result-title"><h2>매뉴얼 미리보기</h2><span class="pw-muted">작성 완료</span></div>', unsafe_allow_html=True)
    st.markdown(markdown)
    html_bytes = make_html(markdown, saved_screens, title)
    col1, col2, col3 = st.columns(3)
    col1.download_button("HTML 다운로드", html_bytes, file_name="사용자_매뉴얼.html", mime="text/html", use_container_width=True)
    try:
        col2.download_button("Word 다운로드", make_docx(markdown, saved_screens, title), file_name="사용자_매뉴얼.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
    except ImportError:
        col2.warning("Word 내보내기 패키지가 없습니다.")
    try:
        col3.download_button("PDF 다운로드", make_pdf(markdown, saved_screens, title), file_name="사용자_매뉴얼.pdf", mime="application/pdf", use_container_width=True)
    except ImportError:
        col3.warning("PDF 내보내기 패키지가 없습니다.")
